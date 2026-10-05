"""Excel (.xlsx) reports: a Summary sheet, the detail sheet (styled headers, filters, frozen
header and name column, real dates/times, totals row with live formulas) and, for period
reports, sheets by location, department and manager."""

import io
from datetime import date, time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.exports.layout import Column, ReportDocument, Table

BRAND = "1F5FAD"
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="E8EEF7")
THIN = Side(style="thin", color="C9D3E3")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

FORMATS = {
    "date": "dd mmm yyyy",
    "time": "hh:mm",
    "minutes": "[h]:mm",  # durations over 24h still show as hours
    "int": "0",
    "percent": "0.0%",
}


def _safe_text(value: str) -> str:
    """Text starting with = + - @ would run as a formula in Excel: prefix it with an apostrophe."""
    return "'" + value if value[:1] in ("=", "+", "-", "@") else value


def _cell_value(value, kind: str):
    if value is None or value == "":
        return None
    if kind == "time":
        hh, mm = str(value).split(":")[:2]
        return time(int(hh), int(mm))
    if kind == "minutes":
        return int(value) / 1440  # Excel stores durations as fractions of a day
    if kind == "date":
        return value if isinstance(value, date) else date.fromisoformat(str(value))
    if kind in ("int", "percent"):
        return value
    return _safe_text(str(value))


def _write_table(ws: Worksheet, table: Table, start_row: int, freeze_name: bool) -> int:
    cols = table.columns
    for c, col in enumerate(cols, start=1):
        cell = ws.cell(row=start_row, column=c, value=col.header)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = max(ws.column_dimensions[get_column_letter(c)].width or 0,
                                                               col.width)
    ws.row_dimensions[start_row].height = 30

    first = start_row + 1
    for r, row in enumerate(table.rows, start=first):
        for c, col in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=c, value=_cell_value(row.get(col.key), col.kind))
            cell.border = BORDER
            if col.kind in FORMATS:
                cell.number_format = FORMATS[col.kind]
                cell.alignment = Alignment(horizontal="center")
    last = first + len(table.rows) - 1

    # Filters on the header row, header (and first column) frozen while scrolling.
    if table.rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{last}"
    ws.freeze_panes = ws.cell(row=first, column=2 if freeze_name else 1)

    if not table.totals:
        return last
    total_row = last + 1
    ws.cell(row=total_row, column=1, value=f"Total ({len(table.rows)})")
    for c, col in enumerate(cols, start=1):
        cell = ws.cell(row=total_row, column=c)
        cell.fill, cell.font, cell.border = TOTAL_FILL, Font(bold=True), BORDER
        letter = get_column_letter(c)
        if col.total and table.rows:
            cell.value = f"=SUM({letter}{first}:{letter}{last})"
            cell.number_format = FORMATS.get(col.kind, "General")
            cell.alignment = Alignment(horizontal="center")
        elif col.kind == "percent" and table.rate_from and table.rows:
            present = get_column_letter(_index(cols, table.rate_from[0]))
            working = get_column_letter(_index(cols, table.rate_from[1]))
            cell.value = f"=IF({working}{total_row}=0,\"\",{present}{total_row}/{working}{total_row})"
            cell.number_format = FORMATS["percent"]
            cell.alignment = Alignment(horizontal="center")
    return total_row


def _index(cols: list[Column], key: str) -> int:
    return next(i for i, c in enumerate(cols, start=1) if c.key == key)


def _summary_sheet(ws: Worksheet, doc: ReportDocument) -> None:
    ws.title = "Summary"
    ws["A1"] = doc.company
    ws["A1"].font = Font(bold=True, size=16, color=BRAND)
    ws["A2"] = doc.title
    ws["A2"].font = Font(bold=True, size=13)
    rows = [("Period", doc.period), ("Generated", doc.generated_label),
            ("Covers", doc.scope)]
    for i, (label, value) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value)
    start = 4 + len(rows) + 1
    ws.cell(row=start, column=1, value="Summary").font = Font(bold=True, size=12)
    for i, (label, value, kind) in enumerate(doc.summary, start=start + 1):
        ws.cell(row=i, column=1, value=label).border = BORDER
        cell = ws.cell(row=i, column=2, value=_cell_value(value, kind))
        cell.border, cell.font = BORDER, Font(bold=True)
        if kind in FORMATS:
            cell.number_format = FORMATS[kind]
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 36


def build_excel(doc: ReportDocument) -> bytes:
    wb = Workbook()
    _summary_sheet(wb.active, doc)
    main = wb.create_sheet(doc.main.title[:31])
    _write_table(main, doc.main, start_row=1, freeze_name=True)
    for table in doc.extra:
        ws = wb.create_sheet(table.title[:31])
        _write_table(ws, table, start_row=1, freeze_name=False)
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        if ws.title != "Summary":
            ws.print_title_rows = "1:1"  # header repeated on every printed page
    wb.properties.title = doc.title
    wb.properties.creator = doc.company
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

