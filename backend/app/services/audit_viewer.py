"""Audit log viewer (HR / Admin), Excel export, and verification of the tamper-evident chain.

Every audit row stores a hash of its own content plus the previous row's hash (database
trigger, Phase 2). Verification recomputes all hashes:
    - a row edited directly in the database      -> its hash no longer matches ("changed")
    - a row deleted from the middle               -> a gap in the sequence numbers ("missing")
    - the newest rows deleted                     -> the last verified row is gone or different
Old rows removed by the retention job are always the oldest ones, which keeps the chain valid.
"""

import io
import json
import uuid
from datetime import date, datetime, time, timedelta
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, or_, select, text
from sqlalchemy.orm import Session, aliased

from app.core import clock
from app.core.errors import AppError
from app.models import AuditCheckpoint, AuditLog, Employee, Organization, User
from app.schemas.security import AuditExportRequest, AuditLogDetail, AuditLogItem, AuditLogPage, ChainResult
from app.services.audit import Actor, audit
from app.services.exports.excel import _safe_text
from app.services.exports.service import ExportFile

MAX_EXPORT_DAYS = 366
MAX_EXPORT_ROWS = 100_000


def _changed(old: dict | None, new: dict | None) -> list[str]:
    old, new = old or {}, new or {}
    return sorted(k for k in set(old) | set(new) if old.get(k) != new.get(k))


def _bounds(db: Session, user: User, start: date | None, end: date | None):
    tz = ZoneInfo(db.get(Organization, user.organization_id).default_timezone)
    lo = datetime.combine(start, time.min, tz) if start else None
    hi = datetime.combine(end + timedelta(days=1), time.min, tz) if end else None
    return lo, hi


def _query(db: Session, user: User, start, end, action, actor_user_id, object_type, object_id, q):
    actor_emp = aliased(Employee)
    query = (
        select(AuditLog, User.email, actor_emp.full_name)
        .outerjoin(User, User.id == AuditLog.actor_user_id)
        .outerjoin(actor_emp, actor_emp.user_id == AuditLog.actor_user_id)
        .where(AuditLog.organization_id == user.organization_id)
    )
    lo, hi = _bounds(db, user, start, end)
    if lo:
        query = query.where(AuditLog.created_at >= lo)
    if hi:
        query = query.where(AuditLog.created_at < hi)
    if action:
        query = query.where(AuditLog.action == action)
    if actor_user_id:
        query = query.where(AuditLog.actor_user_id == actor_user_id)
    if object_type:
        query = query.where(AuditLog.object_type == object_type)
    if object_id:
        query = query.where(AuditLog.object_id == object_id)
    if q:
        like = f"%{q}%"
        query = query.where(or_(AuditLog.action.ilike(like), AuditLog.object_type.ilike(like),
                                AuditLog.object_id.ilike(like), User.email.ilike(like), actor_emp.full_name.ilike(like)))
    return query


def _item(row, email: str | None, name: str | None) -> dict:
    return dict(
        id=row.id, seq=row.seq, created_at=row.created_at, actor=email or "System", actor_name=name,
        actor_role=row.actor_role, action=row.action, object_type=row.object_type, object_id=row.object_id,
        changed=_changed(row.old_value, row.new_value), ip_address=str(row.ip_address) if row.ip_address else None,
    )


def list_logs(db: Session, user: User, *, start=None, end=None, action=None, actor_user_id=None,
              object_type=None, object_id=None, q=None, limit=50, offset=0) -> AuditLogPage:
    query = _query(db, user, start, end, action, actor_user_id, object_type, object_id, q)
    total = db.scalar(select(func.count()).select_from(query.subquery()))
    rows = db.execute(query.order_by(AuditLog.seq.desc()).limit(limit).offset(offset)).all()
    return AuditLogPage(items=[AuditLogItem(**_item(r, e, n)) for r, e, n in rows], total=total)


def detail(db: Session, user: User, log_id: uuid.UUID) -> AuditLogDetail:
    row = db.execute(_query(db, user, None, None, None, None, None, None, None).where(AuditLog.id == log_id)).first()
    if row is None:
        raise AppError(404, "NOT_FOUND", "Audit entry not found.")
    log, email, name = row
    return AuditLogDetail(**_item(log, email, name), old_value=log.old_value, new_value=log.new_value,
                          user_agent=log.user_agent, request_id=log.request_id, row_hash=log.row_hash,
                          prev_hash=log.prev_hash)


def actions(db: Session, user: User) -> list[str]:
    return list(db.scalars(select(AuditLog.action).where(AuditLog.organization_id == user.organization_id)
                           .distinct().order_by(AuditLog.action)))


def export(db: Session, actor: Actor, req: AuditExportRequest) -> ExportFile:
    """Excel file of audit entries. Exporting is itself recorded in the audit log."""
    if req.date_to < req.date_from:
        raise AppError(422, "VALIDATION_ERROR", "'to' must be on or after 'from'.")
    if (req.date_to - req.date_from).days >= MAX_EXPORT_DAYS:
        raise AppError(422, "VALIDATION_ERROR", f"Choose at most {MAX_EXPORT_DAYS} days.")
    user = actor.user
    tz = ZoneInfo(db.get(Organization, user.organization_id).default_timezone)
    query = _query(db, user, req.date_from, req.date_to, req.action, req.actor_user_id, req.object_type, None, req.q)
    rows = db.execute(query.order_by(AuditLog.seq).limit(MAX_EXPORT_ROWS + 1)).all()
    if len(rows) > MAX_EXPORT_ROWS:
        raise AppError(422, "TOO_MANY_ROWS", f"More than {MAX_EXPORT_ROWS:,} entries. Choose a shorter period or a filter.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Audit log"
    headers = ["#", "Time", "Actor", "Name", "Role", "Action", "Object", "Object ID", "Changed fields",
               "Old value", "New value", "IP address", "Request ID"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5FAD")

    def safe(value):
        return _safe_text(value) if isinstance(value, str) else value

    for log, email, name in rows:
        ws.append([safe(v) for v in [
            log.seq, log.created_at.astimezone(tz).replace(tzinfo=None), email or "System", name, log.actor_role,
            log.action, log.object_type, log.object_id, ", ".join(_changed(log.old_value, log.new_value)),
            json.dumps(log.old_value, ensure_ascii=False)[:32000] if log.old_value is not None else None,
            json.dumps(log.new_value, ensure_ascii=False)[:32000] if log.new_value is not None else None,
            str(log.ip_address) if log.ip_address else None, log.request_id,
        ]])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "yyyy-mm-dd hh:mm:ss"
    for letter, width in zip("ABCDEFGHIJKLM", [8, 20, 28, 22, 9, 30, 16, 38, 28, 50, 50, 16, 20]):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buffer = io.BytesIO()
    wb.save(buffer)

    audit(db, actor, "AUDIT_LOG_EXPORTED", "audit_log", None,
          new_value={**req.model_dump(mode="json", exclude_none=True), "rows": len(rows)})
    db.commit()
    name = f"audit-log-{req.date_from.isoformat()}-to-{req.date_to.isoformat()}.xlsx"
    return ExportFile(buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", name)


# --- Chain verification -----------------------------------------------------------------------

_CHAIN_SQL = text("""
WITH c AS (
    SELECT seq, prev_hash, row_hash,
           lag(seq) OVER (ORDER BY seq) AS p_seq,
           lag(row_hash) OVER (ORDER BY seq) AS p_hash,
           audit_log_hash(prev_hash, seq, id, organization_id, actor_user_id, action, object_type,
                          object_id, old_value, new_value, created_at) AS calc
    FROM audit_logs
)
SELECT seq,
       CASE WHEN calc IS DISTINCT FROM row_hash THEN 'changed'
            WHEN p_seq IS NOT NULL AND seq <> p_seq + 1 THEN 'missing'
            ELSE 'relinked' END AS problem
FROM c
WHERE calc IS DISTINCT FROM row_hash
   OR (p_seq IS NOT NULL AND (seq <> p_seq + 1 OR prev_hash IS DISTINCT FROM p_hash))
ORDER BY seq
LIMIT 1
""")

_PROBLEMS = {
    "changed": "Entry #{seq} was changed after it was written.",
    "missing": "Entries before #{seq} were deleted.",
    "relinked": "Entry #{seq} no longer links to the entry before it.",
}


def verify_chain(db: Session) -> ChainResult:
    """Check the whole audit chain and store the result as a checkpoint (caller commits)."""
    now = clock.now()
    stats = db.execute(text("SELECT count(*) AS n, min(seq) AS lo, max(seq) AS hi FROM audit_logs")).one()
    last_hash = db.scalar(text("SELECT row_hash FROM audit_logs WHERE seq = :s"), {"s": stats.hi}) if stats.hi else None
    problem, problem_seq = None, None

    bad = db.execute(_CHAIN_SQL).first()
    if bad is not None:
        problem, problem_seq = _PROBLEMS[bad.problem].format(seq=bad.seq), bad.seq
    else:
        # The newest rows can't be checked by the chain alone: compare with the last good checkpoint.
        previous = db.scalar(select(AuditCheckpoint).where(AuditCheckpoint.ok, AuditCheckpoint.last_seq.is_not(None))
                             .order_by(AuditCheckpoint.created_at.desc()).limit(1))
        if previous is not None and stats.hi is not None and previous.last_seq >= (stats.lo or 0):
            still = db.scalar(text("SELECT row_hash FROM audit_logs WHERE seq = :s"), {"s": previous.last_seq})
            if still is None:
                problem, problem_seq = (f"Entries from #{previous.last_seq} (already verified) were deleted.",
                                        previous.last_seq)
            elif still != previous.last_hash:
                problem, problem_seq = f"Entry #{previous.last_seq} changed since it was verified.", previous.last_seq
        elif previous is not None and stats.hi is None:
            problem, problem_seq = "All audit entries were deleted.", previous.last_seq

    ok = problem is None
    db.add(AuditCheckpoint(ok=ok, last_seq=stats.hi if ok else None, last_hash=last_hash if ok else None,
                           rows_checked=stats.n, problem=problem, problem_seq=problem_seq))
    db.flush()
    return ChainResult(ok=ok, rows_checked=stats.n, last_seq=stats.hi, problem=problem, problem_seq=problem_seq,
                       checked_at=now)


def last_check(db: Session) -> ChainResult | None:
    c = db.scalar(select(AuditCheckpoint).order_by(AuditCheckpoint.created_at.desc()).limit(1))
    if c is None:
        return None
    return ChainResult(ok=c.ok, rows_checked=c.rows_checked, last_seq=c.last_seq, problem=c.problem,
                       problem_seq=c.problem_seq, checked_at=c.created_at)
