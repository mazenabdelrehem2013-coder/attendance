"""Phase 16: audit log viewer, tamper check, alert rules, security dashboard, data retention."""

import io
import random
from datetime import date, datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text

from app.core import clock
from app.db.session import make_session_factory
from app.models import AttendanceEvent, AuditLog, Notification, SecurityAlert, SecurityEvent, User
from app.models.enums import Severity
from app.services import monitoring, retention
from app.services.audit_viewer import verify_chain
from tests.api.attendance_helpers import ann_phone, attempt, set_time  # noqa: F401


@pytest.fixture(scope="module")
def owner_factory():
    return make_session_factory(owner=True, test=True)


def get(client, world, path, who="hr", **params):
    r = client.get(f"/api/v1{path}", headers=world.h(who), params=params)
    assert r.status_code == 200, r.text
    return r.json()


def user_id(session_factory, email):
    with session_factory() as db:
        return db.scalar(select(User.id).where(User.email == email))


def add_events(session_factory, world, event_type, n=1, *, who=None, user=None, ip=None, details=None,
               severity=Severity.HIGH, at=None):
    with session_factory.begin() as db:
        for i in range(n):
            db.add(SecurityEvent(
                organization_id=world.org_id, event_type=event_type, severity=severity,
                employee_id=world.ids[who] if who else None, user_id=user, ip_address=ip,
                details=details or {}, created_at=(at or clock.now()) + timedelta(milliseconds=i),
            ))


def rules_for(session_factory, world, now=None, ip_rule=False):
    """Run the rules for this test company only (the shared test database has many). Failed logins
    without a company (unknown names) count for every company, so other tests' login attempts can
    raise LOGIN_ATTACK_IP here too: it is left out unless asked for."""
    with session_factory() as db:
        new = [a for f in monitoring.evaluate(db, world.org_id, now or clock.now())
               if (a := monitoring.raise_alert(db, world.org_id, f)) is not None]
        db.commit()
        return [(a.rule, a.title) for a in new if ip_rule or a.rule != "LOGIN_ATTACK_IP"]


# --- Audit log viewer -----------------------------------------------------------------------


def test_audit_log_lists_changes_of_this_company_only(client, world):
    client.put(f"/api/v1/employees/{world.ids['ben']}", headers=world.h("hr"), json={"full_name": "Benjamin"})
    page = get(client, world, "/audit-logs", action="EMPLOYEE_UPDATED")
    [entry] = page["items"]
    assert entry["actor"] == world.users["hr"] and entry["actor_name"] == "Hr" and entry["actor_role"] == "HR"
    assert "full_name" in entry["changed"]
    detail = get(client, world, f"/audit-logs/{entry['id']}")
    assert detail["old_value"]["full_name"] == "Ben" and detail["new_value"]["full_name"] == "Benjamin"
    assert len(detail["row_hash"]) == 64
    assert "EMPLOYEE_UPDATED" in get(client, world, "/audit-logs/actions")

    assert get(client, world, "/audit-logs", q=world.users["hr"])["total"] >= 1
    assert get(client, world, "/audit-logs", q="nobody-at-all")["total"] == 0
    for who in ("mgr_a", "ann"):
        assert client.get("/api/v1/audit-logs", headers=world.h(who)).status_code == 403


def test_other_companies_entries_are_invisible(client, world, test_session_factory):
    from tests.api.world import build_world

    other = build_world(test_session_factory, client)
    client.put(f"/api/v1/employees/{other.ids['ann']}", headers=other.h("hr"), json={"full_name": "Other Ann"})
    entry = get(client, other, "/audit-logs", who="hr", action="EMPLOYEE_UPDATED")["items"][0]
    assert client.get(f"/api/v1/audit-logs/{entry['id']}", headers=world.h("hr")).status_code == 404
    assert get(client, world, "/audit-logs", q="Other Ann")["total"] == 0


def test_audit_export_is_safe_and_recorded(client, world):
    client.put(f"/api/v1/employees/{world.ids['ben']}", headers=world.h("hr"), json={"full_name": "=cmd|' /C calc'!A0"})
    today = date.today().isoformat()
    r = client.post("/api/v1/audit-logs/export", headers=world.h("hr"),
                    json={"date_from": "2026-01-01", "date_to": today, "action": "EMPLOYEE_UPDATED"})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content))["Audit log"]
    assert [c.value for c in ws[1]][:6] == ["#", "Time", "Actor", "Name", "Role", "Action"]
    new_value = ws.cell(row=2, column=11).value
    assert "calc" in new_value and not new_value.startswith("=")
    assert get(client, world, "/audit-logs", action="AUDIT_LOG_EXPORTED")["total"] == 1

    too_long = client.post("/api/v1/audit-logs/export", headers=world.h("hr"),
                           json={"date_from": "2024-01-01", "date_to": today})
    assert too_long.status_code == 422


# --- Tamper check ---------------------------------------------------------------------------


def test_chain_verifies(client, world):
    r = client.post("/api/v1/audit-logs/verify", headers=world.h("hr"))
    assert r.status_code == 200 and r.json()["ok"] is True and r.json()["rows_checked"] > 0
    assert get(client, world, "/audit-logs/chain")["ok"] is True


def _middle_entry(db, world):
    """An entry of this company that has a newer entry after it."""
    return db.scalar(select(AuditLog).where(AuditLog.organization_id == world.org_id).order_by(AuditLog.seq).limit(1))


def test_deleted_entry_is_detected(client, world, owner_factory):
    client.put(f"/api/v1/employees/{world.ids['ben']}", headers=world.h("hr"), json={"full_name": "Ben B"})
    client.put(f"/api/v1/employees/{world.ids['ben']}", headers=world.h("hr"), json={"full_name": "Ben BB"})
    with owner_factory() as db:  # an attacker with database access - rolled back afterwards
        victim = _middle_entry(db, world)
        db.execute(text("SET LOCAL app.allow_purge = 'on'"))
        db.execute(text("DELETE FROM audit_logs WHERE seq = :s"), {"s": victim.seq})
        result = verify_chain(db)
        assert result.ok is False and result.problem_seq == victim.seq + 1
        assert "deleted" in result.problem
        db.rollback()


def test_changed_entry_is_detected(client, world, owner_factory):
    client.put(f"/api/v1/employees/{world.ids['ben']}", headers=world.h("hr"), json={"full_name": "Ben C"})
    with owner_factory() as db:
        victim = _middle_entry(db, world)
        db.execute(text("ALTER TABLE audit_logs DISABLE TRIGGER audit_logs_append_only"))
        db.execute(text("UPDATE audit_logs SET action = 'NOTHING_HAPPENED' WHERE seq = :s"), {"s": victim.seq})
        result = verify_chain(db)
        assert result.ok is False and result.problem_seq == victim.seq and "changed" in result.problem
        db.rollback()


def test_deleting_the_newest_entries_is_detected(client, world, owner_factory):
    client.post("/api/v1/audit-logs/verify", headers=world.h("hr"))  # good checkpoint
    with owner_factory() as db:
        newest = db.scalar(select(AuditLog.seq).order_by(AuditLog.seq.desc()).limit(1))
        db.execute(text("SET LOCAL app.allow_purge = 'on'"))
        db.execute(text("DELETE FROM audit_logs WHERE seq = :s"), {"s": newest})
        result = verify_chain(db)
        assert result.ok is False and "already verified" in result.problem
        db.rollback()


# --- Alert rules ----------------------------------------------------------------------------


def test_locked_account_alert_lifecycle(client, world, test_session_factory):
    ann = user_id(test_session_factory, world.users["ann"])
    add_events(test_session_factory, world, "LOGIN_FAILED", user=ann, details={"locked": True}, severity=Severity.MEDIUM)
    assert rules_for(test_session_factory, world) == [("ACCOUNT_LOCKED", f"Account locked: {world.users['ann']}")]
    assert rules_for(test_session_factory, world) == []  # no duplicate

    bell = [i for i in get(client, world, "/notifications")["items"]
            if i["event"] == "SECURITY_ALERT" and "Account locked" in i["title"]]
    assert bell[0]["title"].startswith("Security alert: Account locked") and bell[0]["link"] == "/security"
    assert not [i for i in get(client, world, "/notifications", who="mgr_a")["items"] if i["event"] == "SECURITY_ALERT"]

    def locked(status="active"):
        return [a for a in get(client, world, "/security/alerts", status=status)["items"] if a["rule"] == "ACCOUNT_LOCKED"]

    [alert] = locked()
    assert alert["severity"] == "HIGH" and alert["status"] == "OPEN" and alert["rule_label"] == "Account locked"
    url = f"/api/v1/security/alerts/{alert['id']}"
    assert client.post(url, headers=world.h("hr"), json={"status": "ACKNOWLEDGED"}).json()["status"] == "ACKNOWLEDGED"
    assert client.post(url, headers=world.h("hr"), json={"status": "RESOLVED"}).status_code == 422  # note required
    done = client.post(url, headers=world.h("hr"), json={"status": "RESOLVED", "note": "Ann forgot her password"})
    assert done.json()["status"] == "RESOLVED" and done.json()["handled_by"] == world.users["hr"]
    assert client.post(url, headers=world.h("hr"), json={"status": "RESOLVED", "note": "x"}).status_code == 409
    assert locked() == []
    assert len(locked("RESOLVED")) == 1
    assert get(client, world, "/audit-logs", action="SECURITY_ALERT_UPDATED")["total"] == 2

    assert rules_for(test_session_factory, world) == []  # nothing new since it was resolved
    add_events(test_session_factory, world, "LOGIN_FAILED", user=ann, details={"locked": True},
               at=clock.now() + timedelta(seconds=2))
    assert [r for r, _ in rules_for(test_session_factory, world, clock.now() + timedelta(seconds=3))] == ["ACCOUNT_LOCKED"]


def test_other_rules(client, world, test_session_factory):
    ip = f"203.0.113.{random.randint(1, 254)}"
    add_events(test_session_factory, world, "LOGIN_FAILED", 20, ip=ip, severity=Severity.LOW)
    add_events(test_session_factory, world, "MOCK_LOCATION", 2, who="ann")
    add_events(test_session_factory, world, "INTEGRITY_FAIL", 1, who="ann")
    add_events(test_session_factory, world, "DEVICE_SHARED", 1, who="ben")
    add_events(test_session_factory, world, "REFRESH_TOKEN_REUSE", 1,
               user=user_id(test_session_factory, world.users["cal"]))
    add_events(test_session_factory, world, "OUTSIDE_GEOFENCE", 11, who="cal", severity=Severity.LOW)
    all_found = rules_for(test_session_factory, world, ip_rule=True)
    assert ("LOGIN_ATTACK_IP", f"20 failed logins from {ip} in 15 minutes") in all_found
    found = dict(f for f in all_found if f[0] != "LOGIN_ATTACK_IP")
    assert found["REPEATED_SPOOFING"] == "Ann: 3 serious check-in problems in 24 hours"
    assert found["PHONE_SHARED"] == "Ben tried to use another employee's phone"
    assert found["TOKEN_REUSE"] == f"Possible stolen session: {world.users['cal']}"
    assert found["SUSPICIOUS_SPIKE"] == "15 check-in security events in the last hour"

    items = get(client, world, "/security/alerts")["items"]
    assert [i["severity"] for i in items][:1] == ["HIGH"] and items[-1]["rule"] == "SUSPICIOUS_SPIKE"
    spoof = next(i for i in items if i["rule"] == "REPEATED_SPOOFING")
    assert spoof["employee_name"] == "Ann" and spoof["details"]["types"] == ["INTEGRITY_FAIL", "MOCK_LOCATION"]


def test_below_thresholds_no_alert(client, world, test_session_factory):
    ip = f"198.51.100.{random.randint(1, 254)}"
    add_events(test_session_factory, world, "LOGIN_FAILED", 19, ip=ip)
    add_events(test_session_factory, world, "MOCK_LOCATION", 2, who="ann")
    old = clock.now() - timedelta(hours=30)
    add_events(test_session_factory, world, "MOCK_LOCATION", 5, who="ben", at=old)  # outside the 24-hour window
    found = rules_for(test_session_factory, world, ip_rule=True)
    assert [f for f in found if f[0] != "LOGIN_ATTACK_IP" or ip in f[1]] == []


def test_run_rules_for_all_companies(world, test_session_factory):
    add_events(test_session_factory, world, "DEVICE_SHARED", 1, who="ann")
    assert monitoring.run_rules(test_session_factory) >= 1
    with test_session_factory() as db:
        assert db.scalar(select(SecurityAlert).where(SecurityAlert.organization_id == world.org_id))


# --- Security dashboard ---------------------------------------------------------------------


def test_overview_numbers(client, world, test_session_factory):
    ann = user_id(test_session_factory, world.users["ann"])
    add_events(test_session_factory, world, "MOCK_LOCATION", 2, who="ann")
    add_events(test_session_factory, world, "OUTSIDE_GEOFENCE", 3, who="ben", severity=Severity.LOW)
    add_events(test_session_factory, world, "LOGIN_FAILED", 4, user=ann, ip="192.0.2.77", severity=Severity.LOW)
    add_events(test_session_factory, world, "LOGIN_FAILED", 1, user=ann, ip="192.0.2.77",
               details={"locked": True}, severity=Severity.MEDIUM)
    rules_for(test_session_factory, world)
    o = get(client, world, "/security/overview")
    assert o["total_events"] == 10
    assert o["by_severity"] == {"LOW": 7, "MEDIUM": 1, "HIGH": 2, "CRITICAL": 0}
    assert {t["key"]: t["count"] for t in o["by_type"]} == {"LOGIN_FAILED": 5, "OUTSIDE_GEOFENCE": 3, "MOCK_LOCATION": 2}
    assert o["top_employees"][0]["label"] == "Ann" and o["top_employees"][0]["count"] == 2
    assert o["failed_logins"] >= 5 and o["locked_accounts"] >= 1
    assert {"key": "192.0.2.77", "label": None, "count": 5} in o["top_ips"]
    assert o["open_alerts"]["HIGH"] >= 1  # the locked account
    assert "ACCOUNT_LOCKED" in [a["rule"] for a in get(client, world, "/security/alerts")["items"]]
    assert len(o["per_day"]) == 30 and sum(d["high"] for d in o["per_day"]) == 2
    assert [r["rule"] for r in get(client, world, "/security/rules")][:2] == ["ACCOUNT_LOCKED", "LOGIN_ATTACK_IP"]
    assert client.get("/api/v1/security/overview", headers=world.h("mgr_a")).status_code == 403


# --- Data retention -------------------------------------------------------------------------


def test_retention_settings(client, world):
    items = {i["category"]: i for i in get(client, world, "/retention")}
    assert items["RAW_LOCATION"]["retain_days"] == 365 and items["ATTENDANCE"]["automatic"] is False
    assert items["REPORT_FILES"]["retain_days"] == 90
    assert client.put("/api/v1/retention/RAW_LOCATION", headers=world.h("hr"), json={"retain_days": 90}).status_code == 403
    r = client.put("/api/v1/retention/SECURITY_EVENTS", headers=world.h("admin"), json={"retain_days": 60})
    assert r.status_code == 422 and "at least 90" in r.json()["error"]["message"]
    r = client.put("/api/v1/retention/RAW_LOCATION", headers=world.h("admin"), json={"retain_days": 90})
    assert r.status_code == 200 and r.json()["retain_days"] == 90
    assert get(client, world, "/audit-logs", action="RETENTION_CHANGED")["total"] == 1


def test_retention_clean_up(client, world, ann_phone, set_time, test_session_factory, owner_factory):
    admin = world.h("admin")
    for category, days in (("RAW_LOCATION", 30), ("SECURITY_EVENTS", 90), ("NOTIFICATIONS", 30)):
        assert client.put(f"/api/v1/retention/{category}", headers=admin, json={"retain_days": days}).status_code == 200
    assert attempt(client, world, "ann", ann_phone).json()["result"] == "ACCEPTED"  # Monday 5 Oct 2026
    add_events(test_session_factory, world, "MOCK_LOCATION", 1, who="ann", at=clock.now())
    with test_session_factory() as db:
        assert db.scalar(select(AttendanceEvent.latitude).where(AttendanceEvent.employee_id == world.ids["ann"]))

    later = datetime(2027, 3, 1, tzinfo=timezone.utc)  # > 90 days later
    results = retention.purge(owner_factory, later)[str(world.org_id)]
    assert results["RAW_LOCATION"] == 1 and results["SECURITY_EVENTS"] >= 1 and results["NOTIFICATIONS"] >= 1
    with test_session_factory() as db:
        event = db.scalar(select(AttendanceEvent).where(AttendanceEvent.employee_id == world.ids["ann"]))
        assert (event.latitude, event.longitude, event.accuracy_m) == (None, None, None)
        assert event.distance_m is not None and event.result.value == "ACCEPTED"  # the rest stays
        assert not db.scalar(select(SecurityEvent).where(SecurityEvent.organization_id == world.org_id))
        hr_id = db.scalar(select(User.id).where(User.email == world.users["hr"]))
        assert not db.scalar(select(Notification).where(Notification.user_id == hr_id))
    assert get(client, world, "/retention/last-run")["results"]["RAW_LOCATION"] == 1


def test_only_gps_can_be_cleared_and_only_with_the_switch(owner_factory, world, client, ann_phone):
    attempt(client, world, "ann", ann_phone)
    with owner_factory() as db:
        db.execute(text("SET LOCAL app.allow_redact = 'on'"))
        with pytest.raises(Exception, match="append-only"):
            with db.begin_nested():  # anything besides the coordinates: refused
                db.execute(text("UPDATE attendance_events SET latitude = NULL, longitude = NULL, accuracy_m = NULL, "
                                "result = 'REJECTED' WHERE employee_id = :e"), {"e": world.ids["ann"]})
        db.rollback()
    with owner_factory() as db:
        with pytest.raises(Exception, match="append-only"):  # without the switch: refused
            db.execute(text("UPDATE attendance_events SET latitude = NULL, longitude = NULL, accuracy_m = NULL "
                            "WHERE employee_id = :e"), {"e": world.ids["ann"]})
        db.rollback()
    app = make_session_factory(test=True)
    with app() as db:
        with pytest.raises(Exception, match="permission denied"):  # the API's own role: never
            db.execute(text("SET LOCAL app.allow_redact = 'on'"))
            db.execute(text("UPDATE attendance_events SET latitude = NULL WHERE employee_id = :e"),
                       {"e": world.ids["ann"]})
        db.rollback()




def test_scheduler_runs_the_nightly_checks_once(world, set_time, test_session_factory):
    from app.jobs.scheduler import tick

    set_time(1, 30, day=(2026, 10, 6))
    first = tick(test_session_factory)
    assert first.chain_ok is None and first.retention_ran is False  # before 02:00
    set_time(2, 30, day=(2026, 10, 6))
    night = tick(test_session_factory)
    assert night.chain_ok is True and night.retention_ran is True
    assert tick(test_session_factory).retention_ran is False  # once per night


def test_the_gps_switch_does_not_open_other_tables(owner_factory):
    with owner_factory() as db:
        db.execute(text("SET LOCAL app.allow_redact = 'on'"))
        for sql in ("UPDATE audit_logs SET action = 'HACKED'", "UPDATE security_events SET severity = 'LOW'"):
            with pytest.raises(Exception, match="append-only"):
                with db.begin_nested():
                    db.execute(text(sql))
        db.rollback()
