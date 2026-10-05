"""Phase 15 end to end: bell alerts, scheduled reports saved as files ("Ready reports"),
downloading them (dashboard + app use the same endpoints), clean-up, and the scheduler."""

import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text, update

from app.core import clock
from app.db.session import make_session_factory
from app.jobs.close_day import close_day
from app.jobs.scheduler import _LOCK_ID, tick
from app.models import AuditLog, Organization, ReportFile, ReportSetting
from app.models.enums import ReportType
from app.services import report_schedules
from tests.api.attendance_helpers import ann_phone, approved_device, attempt, set_time  # noqa: F401


def alerts(client, world, who):
    r = client.get("/api/v1/notifications", headers=world.h(who))
    assert r.status_code == 200, r.text
    return r.json()


def titles(client, world, who, event=None):
    return [i["title"] for i in alerts(client, world, who)["items"] if event is None or i["event"] == event]


def ready(client, world, who):
    r = client.get("/api/v1/reports/ready", headers=world.h(who))
    assert r.status_code == 200, r.text
    return r.json()


def download(client, world, who, file_id):
    return client.get(f"/api/v1/reports/ready/{file_id}/download", headers=world.h(who))


def names_in(content: bytes, sheet="Employees") -> list[str]:
    ws = load_workbook(io.BytesIO(content))[sheet]
    return [r[0].value for r in ws.iter_rows(min_row=2) if r[0].value and not str(r[0].value).startswith("Total")]


# --- Bell alerts ----------------------------------------------------------------------------


def test_new_phone_alerts_hr(client, world):
    approved_device(client, world, "ann")
    hr = alerts(client, world, "hr")
    assert hr["unread"] == 1 and hr["items"][0]["title"] == "New phone to approve: Ann"
    assert hr["items"][0]["link"] == "/phones" and "Pixel 8" in hr["items"][0]["body"]
    assert alerts(client, world, "ann")["unread"] == 0  # nobody is told about themselves


def test_flagged_check_in_alerts_hr_with_the_reason(client, world, ann_phone):
    assert attempt(client, world, "ann", ann_phone, is_mock_location=True).json()["result"] == "FLAGGED"
    flag = next(i for i in alerts(client, world, "hr")["items"] if i["event"] == "SUSPICIOUS_CHECK_IN")
    assert flag["title"] == "Check-in needs review: Ann" and flag["link"] == "/review"
    assert "Fake-GPS app detected" in flag["body"] and "Lagos" in flag["body"]
    assert not titles(client, world, "mgr_a", "SUSPICIOUS_CHECK_IN")


def test_late_check_in_alerts_the_manager_once(client, world, ann_phone, set_time):
    set_time(9, 30)
    assert attempt(client, world, "ann", ann_phone).json()["result"] == "ACCEPTED"
    attempt(client, world, "ann", ann_phone, action="CHECK_OUT")
    set_time(10, 0)
    attempt(client, world, "ann", ann_phone)  # second visit the same day: no second alert
    assert titles(client, world, "mgr_a", "LATE_EMPLOYEE") == ["Late arrival: Ann"]
    assert not titles(client, world, "mgr_b", "LATE_EMPLOYEE")


def test_hr_changes_who_gets_an_alert(client, world, ann_phone, set_time):
    settings = client.get("/api/v1/hr/alert-settings", headers=world.h("hr")).json()
    assert {s["event"] for s in settings} == {"SUSPICIOUS_CHECK_IN", "REJECTED_CHECK_IN", "DEVICE_APPROVAL_REQUESTED",
                                              "LATE_EMPLOYEE", "MISSING_CHECKOUT", "HIGH_ABSENCE", "REPORT_GENERATED",
                                              "SECURITY_ALERT"}
    r = client.put("/api/v1/hr/alert-settings/LATE_EMPLOYEE", headers=world.h("hr"),
                   json={"enabled": True, "roles": ["HR"]})
    assert r.status_code == 200 and r.json()["roles"] == ["HR"]
    set_time(9, 40)
    attempt(client, world, "ann", ann_phone)
    assert titles(client, world, "hr", "LATE_EMPLOYEE") == ["Late arrival: Ann"]
    assert not titles(client, world, "mgr_a", "LATE_EMPLOYEE")

    client.put("/api/v1/hr/alert-settings/DEVICE_APPROVAL_REQUESTED", headers=world.h("hr"),
               json={"enabled": False, "roles": ["HR"]})
    approved_device(client, world, "ben")
    assert "New phone to approve: Ben" not in titles(client, world, "hr", "DEVICE_APPROVAL_REQUESTED")


def test_reading_alerts(client, world):
    approved_device(client, world, "ann")
    approved_device(client, world, "ben")
    items = alerts(client, world, "hr")["items"]
    assert len(items) == 2
    client.post(f"/api/v1/notifications/{items[0]['id']}/read", headers=world.h("ann"))  # not hers: no effect
    assert alerts(client, world, "hr")["unread"] == 2
    client.post(f"/api/v1/notifications/{items[0]['id']}/read", headers=world.h("hr"))
    assert alerts(client, world, "hr")["unread"] == 1
    client.post("/api/v1/notifications/read-all", headers=world.h("hr"))
    assert alerts(client, world, "hr")["unread"] == 0


def test_end_of_day_alerts(client, world, ann_phone, set_time, test_session_factory):
    client.put("/api/v1/hr/alert-settings/HIGH_ABSENCE", headers=world.h("hr"),
               json={"enabled": True, "roles": ["MANAGER", "HR"], "thresholds": {"absences": 1, "days": 30}})
    attempt(client, world, "ann", ann_phone)  # Monday: checks in, never checks out; Ben never comes
    set_time(1, 0, day=(2026, 10, 6))
    for _ in range(2):  # running the job twice doesn't repeat alerts
        with test_session_factory() as db:
            close_day(db, world.org_id, date(2026, 10, 5))
    got = titles(client, world, "mgr_a")
    assert got.count("Missing check-out: Ann") == 1 and got.count("Frequent absence: Ben") == 1
    assert "Frequent absence: Ben" in titles(client, world, "hr")


def test_only_hr_manages_schedules_and_alerts(client, world):
    for path in ("/hr/alert-settings", "/hr/report-schedules"):
        assert client.get(f"/api/v1{path}", headers=world.h("mgr_a")).status_code == 403
    assert client.get("/api/v1/reports/ready", headers=world.h("ann")).status_code == 403


# --- Scheduled reports ----------------------------------------------------------------------

WEEKLY = {"name": "Weekly for HR and managers", "report": "weekly", "frequency": "weekly", "time": "08:00",
          "weekday": 0, "formats": ["EXCEL", "PDF"], "recipient_roles": ["HR", "MANAGER"]}


def test_schedule_crud(client, world, set_time):
    hr = world.h("hr")
    r = client.post("/api/v1/hr/report-schedules", headers=hr, json=WEEKLY)
    assert r.status_code == 201, r.text
    s = r.json()
    assert s["description"] == "Every Monday at 08:00" and s["timezone"] == "Africa/Lagos"
    assert s["next_run_at"].startswith("2026-10-12T08:00:00")  # set_time: Monday 5 Oct 09:00
    assert client.post("/api/v1/hr/report-schedules", headers=hr, json={**WEEKLY, "frequency": "daily"}).status_code == 422
    assert client.post("/api/v1/hr/report-schedules", headers=hr, json={**WEEKLY, "recipient_roles": []}).status_code == 422

    changed = client.put(f"/api/v1/hr/report-schedules/{s['id']}", headers=hr,
                         json={**WEEKLY, "time": "07:30", "is_enabled": False}).json()
    assert changed["description"] == "Every Monday at 07:30" and changed["next_run_at"] is None
    assert client.delete(f"/api/v1/hr/report-schedules/{s['id']}", headers=hr).status_code == 204
    assert client.get("/api/v1/hr/report-schedules", headers=hr).json() == []


def test_run_now_creates_company_and_team_files(client, world, set_time, test_session_factory):
    s = client.post("/api/v1/hr/report-schedules", headers=world.h("hr"), json=WEEKLY).json()
    r = client.post(f"/api/v1/hr/report-schedules/{s['id']}/run-now", headers=world.h("hr"))
    assert r.json()["files_created"] == 6  # (company + 2 managers) x (Excel + PDF)

    hr_files = ready(client, world, "hr")
    assert hr_files["total"] == 2 and hr_files["keep_days"] == 90
    xlsx = next(f for f in hr_files["items"] if f["format"] == "EXCEL")
    assert xlsx["title"] == "Weekly attendance report" and xlsx["scope"] == "All employees"
    assert xlsx["period"] == "28 September 2026 – 04 October 2026"
    assert xlsx["filename"] == "attendance-weekly-2026-09-28-to-2026-10-04.xlsx"
    assert xlsx["schedule_name"] == "Weekly for HR and managers"

    got = download(client, world, "hr", xlsx["id"])
    assert got.status_code == 200 and got.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="attendance-weekly-2026-09-28-to-2026-10-04.xlsx"' in got.headers["content-disposition"]
    assert set(names_in(got.content)) == {"Ann", "Ben", "Cal", "Hr", "Mgr_A", "Mgr_B"}

    # Each manager sees only their own team's files.
    mgr_a = ready(client, world, "mgr_a")["items"]
    assert {f["scope"] for f in mgr_a} == {"Your team"} and len(mgr_a) == 2
    team_xlsx = next(f for f in mgr_a if f["format"] == "EXCEL")
    assert names_in(download(client, world, "mgr_a", team_xlsx["id"]).content) == ["Ann", "Ben"]
    assert download(client, world, "mgr_b", team_xlsx["id"]).status_code == 404
    assert download(client, world, "mgr_a", xlsx["id"]).status_code == 404  # not the company-wide file

    # People are told, and the download is recorded.
    assert "Report ready: Weekly attendance report" in titles(client, world, "mgr_b", "REPORT_GENERATED")
    with test_session_factory() as db:
        assert db.scalar(select(AuditLog).where(AuditLog.organization_id == world.org_id,
                                                AuditLog.action == "REPORT_DOWNLOADED"))


def test_schedule_runs_once_at_its_time(client, world, set_time, test_session_factory):
    daily = {"name": "Morning", "report": "daily", "frequency": "daily", "time": "08:00",
             "days": [0, 1, 2, 3, 4, 5], "formats": ["PDF"], "recipient_roles": ["HR"]}
    client.post("/api/v1/hr/report-schedules", headers=world.h("hr"), json=daily)

    set_time(7, 59, day=(2026, 10, 6))
    report_schedules.run_due(test_session_factory)
    assert ready(client, world, "hr")["total"] == 0

    set_time(8, 1, day=(2026, 10, 6))
    report_schedules.run_due(test_session_factory)
    report_schedules.run_due(test_session_factory)  # again: nothing new
    [f] = ready(client, world, "hr")["items"]
    assert f["period"] == "Monday 05 October 2026"  # the previous working day

    set_time(8, 0, day=(2026, 10, 11))  # Sunday: not a report day
    report_schedules.run_due(test_session_factory)
    set_time(8, 5, day=(2026, 10, 12))  # Monday: reports on Saturday, skipping Sunday
    report_schedules.run_due(test_session_factory)
    periods = [f["period"] for f in ready(client, world, "hr")["items"]]
    assert periods == ["Saturday 10 October 2026", "Monday 05 October 2026"]


def test_missed_time_is_skipped(client, world, set_time, test_session_factory):
    client.post("/api/v1/hr/report-schedules", headers=world.h("hr"), json={**WEEKLY, "recipient_roles": ["HR"]})
    set_time(20, 0, day=(2026, 10, 12))  # 12 hours after Monday 08:00 (server was down)
    report_schedules.run_due(test_session_factory)
    assert ready(client, world, "hr")["total"] == 0


def test_old_files_are_deleted(client, world, set_time, test_session_factory):
    s = client.post("/api/v1/hr/report-schedules", headers=world.h("hr"), json={**WEEKLY, "recipient_roles": ["HR"]}).json()
    client.post(f"/api/v1/hr/report-schedules/{s['id']}/run-now", headers=world.h("hr"))
    [old, keep] = ready(client, world, "hr")["items"]
    with make_session_factory(owner=True, test=True).begin() as db:  # pretend one file is 91 days old
        db.execute(update(ReportFile).where(ReportFile.id == old["id"])
                   .values(created_at=clock.now() - timedelta(days=91)))
    assert report_schedules.purge_old_files(test_session_factory) >= 1
    assert [f["id"] for f in ready(client, world, "hr")["items"]] == [keep["id"]]
    assert download(client, world, "hr", old["id"]).status_code == 404


def test_schedules_saved_with_only_a_cron_expression_still_load(client, world, test_session_factory):
    with test_session_factory.begin() as db:  # how the development seed stores them
        db.add(ReportSetting(organization_id=world.org_id, name="Seeded daily", report_type=ReportType.DAILY,
                             cron_expression="30 9 * * 1-6", formats=["PDF"], recipient_roles=["HR"]))
    [s] = client.get("/api/v1/hr/report-schedules", headers=world.h("hr")).json()
    assert (s["frequency"], s["days"], s["time"]) == ("daily", [0, 1, 2, 3, 4, 5], "09:30")
    assert s["description"] == "Monday to Saturday at 09:30"


# --- Scheduler ------------------------------------------------------------------------------


def test_scheduler_skips_when_another_one_is_running(test_session_factory):
    with test_session_factory() as other:
        assert other.scalar(text("SELECT pg_try_advisory_lock(:k)"), {"k": _LOCK_ID})
        try:
            assert tick(test_session_factory).ran is False
        finally:
            other.execute(text("SELECT pg_advisory_unlock(:k)"), {"k": _LOCK_ID})


def test_scheduler_closes_yesterday_once(world, set_time, test_session_factory):
    with test_session_factory() as db:
        name = f"{db.get(Organization, world.org_id).name} 2026-10-05"
    set_time(0, 10, day=(2026, 10, 6))
    assert name not in tick(test_session_factory).days_closed  # before 00:15
    set_time(0, 20, day=(2026, 10, 6))
    assert name in tick(test_session_factory).days_closed
    assert name not in tick(test_session_factory).days_closed
