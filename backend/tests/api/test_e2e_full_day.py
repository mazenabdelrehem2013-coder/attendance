"""Phase 17 end-to-end: one complete working day through the whole system, the way people
really use it - HR adds a new employee, the employee sets a password and registers a phone,
HR approves it, people check in (one late, one with a fake-GPS app), HR reviews, people check
out, the night job closes the day, managers/HR see the reports and download them, alerts
arrive, and the audit trail is complete and intact."""

import io
from datetime import date, datetime, timezone

from openpyxl import load_workbook

from app.services import monitoring
from tests.api.attendance_helpers import approved_device, attempt, register, set_time  # noqa: F401

NEW_PASSWORD = "Dayo-Strong-Pass-2026"


def titles(client, world, who, event):
    r = client.get("/api/v1/notifications", headers=world.h(who))
    return [i["title"] for i in r.json()["items"] if i["event"] == event]


def test_a_full_working_day(client, world, set_time, test_session_factory):
    hr = world.h("hr")
    set_time(8, 0)  # Monday 5 Oct 2026, Lagos

    # 1. HR adds a new employee in Lagos, managed by mgr_a.
    r = client.post("/api/v1/employees", headers=hr, json={
        "full_name": "Dayo Ade", "employee_code": "DAYO-1", "email": f"dayo-{world.org_id.hex[:6]}@example.com",
        "manager_id": str(world.manager_ids["mgr_a"]), "department_id": str(world.dept_id),
        "location_ids": [str(world.lagos_id)], "hire_date": "2026-10-01",
    })
    assert r.status_code == 201, r.text
    dayo_id, temporary = r.json()["employee"]["id"], r.json()["temporary_password"]

    # 2. Dayo logs in with the temporary password and must change it before anything else.
    login = client.post("/api/v1/auth/login", json={"identifier": "DAYO-1", "password": temporary,
                                                     "client_type": "MOBILE"}).json()
    assert login["user"]["must_change_password"] is True
    blocked = client.get("/api/v1/attendance/today", headers={"Authorization": f"Bearer {login['access_token']}"})
    assert blocked.status_code == 403 and blocked.json()["error"]["code"] == "PASSWORD_CHANGE_REQUIRED"
    changed = client.post("/api/v1/auth/change-password", headers={"Authorization": f"Bearer {login['access_token']}"},
                          json={"current_password": temporary, "new_password": NEW_PASSWORD})
    assert changed.status_code == 200, changed.text
    world.tokens["dayo"], world.ids["dayo"] = changed.json()["access_token"], dayo_id

    # 3. Dayo registers a phone; it can't be used until HR approves it.
    phone = register(client, world, "dayo")
    assert titles(client, world, "hr", "DEVICE_APPROVAL_REQUESTED") == ["New phone to approve: Dayo Ade"]
    early = client.post("/api/v1/attendance/challenge", headers=world.h("dayo"),
                        json={"action": "CHECK_IN", "device_id": phone["id"]})
    assert early.status_code == 403 and early.json()["error"]["code"] == "DEVICE_PENDING_APPROVAL"
    assert client.post(f"/api/v1/devices/{phone['id']}/approve", headers=hr).status_code == 200

    # 4. Morning: Ann on time but with a fake-GPS app, Dayo late, Ben never comes.
    ann_phone = approved_device(client, world, "ann")
    set_time(9, 5)
    assert attempt(client, world, "ann", ann_phone, is_mock_location=True).json()["result"] == "FLAGGED"
    set_time(9, 25)
    dayo_in = attempt(client, world, "dayo", phone["id"]).json()
    assert dayo_in["result"] == "ACCEPTED"
    assert titles(client, world, "mgr_a", "LATE_EMPLOYEE") == ["Late arrival: Dayo Ade"]
    assert "Check-in needs review: Ann" in titles(client, world, "hr", "SUSPICIOUS_CHECK_IN")

    # 5. HR reviews Ann's flagged check-in and accepts it (she was at the office; the app is removed).
    [item] = client.get("/api/v1/hr/review", headers=hr, params={"state": "PENDING"}).json()["items"]
    assert item["employee_name"] == "Ann" and item["reason_code"] == "MOCK_LOCATION"
    reviewed = client.post(f"/api/v1/hr/review/{item['event_id']}", headers=hr,
                           json={"decision": "APPROVED", "note": "Checked CCTV; fake-GPS app removed"})
    assert reviewed.status_code == 200, reviewed.text

    # 6. Afternoon: Dayo leaves early; Ann forgets to check out.
    set_time(17, 0)
    assert attempt(client, world, "dayo", phone["id"], action="CHECK_OUT").json()["result"] == "ACCEPTED"

    # 7. Night: the scheduler closes Monday.
    from app.jobs.scheduler import tick

    set_time(0, 30, day=(2026, 10, 6))
    assert any(d.endswith("2026-10-05") for d in tick(test_session_factory).days_closed)
    assert "Missing check-out: Ann" in titles(client, world, "mgr_a", "MISSING_CHECKOUT")

    # 8. Tuesday morning: the manager and HR look at Monday.
    set_time(9, 0, day=(2026, 10, 6))
    team = client.get("/api/v1/reports/daily", headers=world.h("mgr_a"), params={"date": "2026-10-05"}).json()
    rows = {r["employee"]: r for r in team["rows"]}
    assert set(rows) == {"Ann", "Ben", "Dayo Ade"}  # only mgr_a's team
    assert (rows["Dayo Ade"]["status"], rows["Dayo Ade"]["departure_status"]) == ("LATE", "EARLY_DEPARTURE")
    assert rows["Dayo Ade"]["check_in"] == "09:25" and rows["Dayo Ade"]["worked_minutes"] == 455
    assert rows["Ben"]["status"] == "ABSENT"
    assert rows["Ann"]["departure_status"] == "MISSING_CHECKOUT"

    excel = client.post("/api/v1/reports/export/excel", headers=hr, json={"report": "daily", "date": "2026-10-05"})
    ws = load_workbook(io.BytesIO(excel.content))["Attendance"]
    status = {r[0].value: r[8].value for r in ws.iter_rows(min_row=2) if r[0].value}
    assert status["Dayo Ade"] == "Late · Left early" and status["Ben"] == "Absent"

    # Dayo's own history shows the same day.
    [day] = client.get("/api/v1/attendance/history", headers=world.h("dayo"),
                       params={"from": "2026-10-05", "to": "2026-10-05"}).json()
    assert day["date"] == "2026-10-05" and day["worked_minutes"] == 455

    # 9. Security: one fake-GPS event on the dashboard; no alert rule fires for a single event.
    # (Security events carry the database's real time, the test clock is simulated: last 30 days.)
    overview = client.get("/api/v1/security/overview", headers=hr).json()
    assert {t["key"]: t["count"] for t in overview["by_type"]}.get("MOCK_LOCATION") == 1
    with test_session_factory() as db:
        assert [f.rule for f in monitoring.evaluate(db, world.org_id, datetime.now(timezone.utc))
                if f.rule != "LOGIN_ATTACK_IP"] == []

    # 10. The audit trail has every step, and nothing in it was altered.
    actions = {i["action"] for i in client.get("/api/v1/audit-logs", headers=hr, params={"limit": 200}).json()["items"]}
    assert {"EMPLOYEE_CREATED", "DEVICE_REGISTRATION_REQUESTED", "DEVICE_APPROVED", "ATTENDANCE_EVENT_REVIEWED",
            "REPORT_EXPORTED"} <= actions
    assert client.post("/api/v1/audit-logs/verify", headers=hr).json()["ok"] is True
    assert date(2026, 10, 5)  # the day under test
